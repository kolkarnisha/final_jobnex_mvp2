"""Job-description extraction and document text extraction services."""
from __future__ import annotations

import io
import re
import zipfile
from collections import Counter
from pathlib import Path
from typing import Any

try:
    from analyzer import CERTIFICATIONS, SKILLS, _extract, normalize_text
except ImportError:  # package import
    from .analyzer import CERTIFICATIONS, SKILLS, _extract, normalize_text

MAX_JOB_TEXT_LENGTH = 200_000
MAX_UPLOAD_SIZE = 10 * 1024 * 1024
ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt"}

_STOP_WORDS = set(
    "a an and are as at be by can candidates candidate collaborate for from has have "
    "in into is it its join of on or our required should team the their this to with "
    "will work years experience".split()
)


def extract_job_details(text: str, file_name: str | None = None) -> dict[str, Any]:
    """Extract consistent, UI-shaped job details from plain-text job content."""
    if not isinstance(text, str) or not text.strip():
        raise ValueError("A non-empty job description is required")
    if len(text) > MAX_JOB_TEXT_LENGTH:
        raise ValueError(f"Job description must be at most {MAX_JOB_TEXT_LENGTH} characters")

    clean_text = text.strip()
    lines = [line.strip() for line in clean_text.splitlines() if line.strip()]
    lower_text = normalize_text(clean_text)

    title_match = re.search(
        r"(?im)^\s*(?:job\s*)?(?:title|role|position)\s*[:\-]\s*([^\r\n,]{2,120})",
        clean_text,
    )
    title = title_match.group(1).strip() if title_match else ""
    if not title and lines and len(lines[0]) <= 120:
        title = re.sub(r"^#+\s*", "", lines[0]).strip()
    if not title:
        title = "Software Engineer"

    company_match = re.search(
        r"(?im)^\s*(?:company|organization|organisation|employer|client)\s*[:\-]\s*([^\r\n,]{2,120})",
        clean_text,
    )
    company = company_match.group(1).strip() if company_match else None

    location_match = re.search(r"(?im)^\s*location\s*[:\-]\s*([^\r\n]{2,160})", clean_text)
    location = location_match.group(1).strip() if location_match else None

    experience_match = re.search(
        r"(?i)\b\d{1,2}\s*(?:\+|to|-)?\s*\d{0,2}\s*(?:years?|yrs?)(?:\s+of)?(?:\s+experience)?",
        clean_text,
    )
    experience = experience_match.group(0).strip() if experience_match else None

    preferred_match = re.search(
        r"(?i)\b(preferred|nice\s+to\s+have|bonus|desired qualifications|preferred qualifications)\b",
        clean_text,
    )
    all_skills = _extract(lower_text, SKILLS)
    if preferred_match:
        required_text = normalize_text(clean_text[: preferred_match.start()])
        preferred_text = normalize_text(clean_text[preferred_match.start() :])
        required_skills = _extract(required_text, SKILLS)
        preferred_skills = _extract(preferred_text, SKILLS)
        unclassified = set(all_skills) - set(required_skills) - set(preferred_skills)
        required_skills = sorted(set(required_skills) | unclassified)
    else:
        required_skills = all_skills
        preferred_skills = []

    words = re.findall(r"[a-z][a-z0-9+#.]{2,}", lower_text)
    frequencies = Counter(word.rstrip(".") for word in words if word.rstrip(".") not in _STOP_WORDS)
    keywords = [
        word for word, _ in frequencies.most_common(24)
        if word not in _STOP_WORDS
    ]

    responsibilities: list[str] = []
    in_responsibilities = False
    for line in lines:
        if re.search(r"(?i)\b(responsibilities|what you(?:'|’)ll do|what you will do|duties)\b", line):
            in_responsibilities = True
            continue
        if in_responsibilities and re.search(
            r"(?i)\b(requirements|qualifications|preferred|benefits|about us)\b", line
        ):
            in_responsibilities = False
        if in_responsibilities:
            item = re.sub(r"^\s*(?:[-*•]|\d+[.)])\s*", "", line).strip()
            if len(item) >= 20:
                responsibilities.append(item[:500])
    if not responsibilities:
        responsibilities = [
            re.sub(r"^\s*(?:[-*•]|\d+[.)])\s*", "", line).strip()
            for line in lines
            if re.match(r"^\s*(?:[-*•]|\d+[.)])\s+", line)
            and len(line.strip()) >= 20
        ][:8]

    certs = _extract(lower_text, CERTIFICATIONS)
    job_type_match = re.search(
        r"(?i)\b(full[\s-]?time|part[\s-]?time|contract|temporary|internship)\b",
        clean_text,
    )
    return {
        "title": title[:200],
        "company": company[:200] if company else None,
        "location": location[:200] if location else None,
        "jobType": job_type_match.group(1) if job_type_match else None,
        "experience": experience,
        "requiredSkills": required_skills,
        "preferredSkills": preferred_skills,
        "certifications": certs,
        "keywords": keywords,
        "responsibilities": responsibilities,
        "rawText": clean_text,
        "sourceFile": Path(file_name).name if file_name else None,
    }


def extract_uploaded_text(file_name: str, content: bytes) -> str:
    """Extract UTF-8 text from supported document formats with bounded input size."""
    extension = Path(file_name).suffix.lower()
    if extension not in ALLOWED_EXTENSIONS:
        raise ValueError("Only PDF, DOCX, and TXT job descriptions are supported")
    if not content:
        raise ValueError("The uploaded file is empty")
    if len(content) > MAX_UPLOAD_SIZE:
        raise ValueError("Uploaded file must be 10 MB or smaller")

    if extension == ".txt":
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError:
            try:
                text = content.decode("utf-16")
            except UnicodeDecodeError as exc:
                raise ValueError("Text file must use UTF-8 or UTF-16 encoding") from exc
    elif extension == ".pdf":
        if not content.startswith(b"%PDF-"):
            raise ValueError("The uploaded file is not a valid PDF document")
        try:
            from pypdf import PdfReader
        except ImportError as exc:  # pragma: no cover - dependency is declared
            raise RuntimeError("PDF support is unavailable; install the pypdf dependency") from exc
        try:
            reader = PdfReader(io.BytesIO(content), strict=True)
            if reader.is_encrypted:
                raise ValueError("Encrypted PDF files are not supported")
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
        except ValueError:
            raise
        except Exception as exc:
            raise ValueError("Could not read the uploaded PDF document") from exc
    else:
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                expanded_size = sum(info.file_size for info in archive.infolist())
                if expanded_size > MAX_UPLOAD_SIZE * 20:
                    raise ValueError("DOCX contents exceed the supported size limit")
                if "word/document.xml" not in archive.namelist():
                    raise ValueError("The uploaded DOCX document is invalid")
            from docx import Document
        except ValueError:
            raise
        except (zipfile.BadZipFile, KeyError) as exc:
            raise ValueError("The uploaded DOCX document is invalid") from exc
        except ImportError as exc:  # pragma: no cover - dependency is declared
            raise RuntimeError("DOCX support is unavailable; install the python-docx dependency") from exc
        try:
            document = Document(io.BytesIO(content))
            text = "\n".join(paragraph.text for paragraph in document.paragraphs)
            for table in document.tables:
                text += "\n" + "\n".join(
                    " ".join(cell.text for cell in row.cells) for row in table.rows
                )
        except Exception as exc:
            raise ValueError("Could not read the uploaded DOCX document") from exc

    if not text.strip():
        raise ValueError("No readable job-description text was found in the uploaded file")
    if len(text) > MAX_JOB_TEXT_LENGTH:
        raise ValueError(f"Extracted job description must be at most {MAX_JOB_TEXT_LENGTH} characters")
    return text
